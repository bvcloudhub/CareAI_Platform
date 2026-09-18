"""Dynamic, role-scoped reports for CareAI."""
from __future__ import annotations

import csv
import io
import json
from datetime import date, datetime, timedelta
from html import escape
from typing import Any, Callable

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from services.access_scope import accessible_patient_ids
from services.db import query_db
from services.population_service import normalise_population_filters, population_dashboard, population_filter_options

try:
    from openpyxl import Workbook
except Exception:  # pragma: no cover - handled by UI capability flag
    Workbook = None

REPORT_DEFINITIONS = {
    "patient_summary": {
        "label": "Patient summary",
        "description": "Current patient master, conditions, risk, assignment, latest activity and device status.",
    },
    "high_critical_risk": {
        "label": "High / Critical risk patients",
        "description": "Patients whose current CareAI priority is High or Critical.",
    },
    "vitals_summary": {
        "label": "Vitals summary",
        "description": "Vital measurements recorded in the selected period.",
    },
    "agentic_events": {
        "label": "Agentic Care events",
        "description": "Agentic Care runs, severity, status and workflow timing.",
    },
    "fall_deterioration": {
        "label": "Fall / deterioration events",
        "description": "Recorded fall and deterioration signals from Agentic Care, care events and device events.",
    },
    "consultations": {
        "label": "Consultations / teleconsultations",
        "description": "Teleconsultation requests and session status.",
    },
    "family_notifications": {
        "label": "Family notifications",
        "description": "Notifications addressed to family/delegated contacts.",
    },
    "device_measurements": {
        "label": "Wearable / device measurements",
        "description": "Measurements received from CareAI device adapters and wearable integrations.",
    },
    "outcome_metrics": {
        "label": "Outcomes and response-time metrics",
        "description": "Agentic workflow response time, intervention and outcome metrics.",
    },
    "clinical_alerts": {
        "label": "Clinical alerts",
        "description": "Clinical alerts, severity and current status.",
    },
    "population_summary": {
        "label": "Population-level summary",
        "description": "Population KPIs and cohort counts for the selected patient scope.",
    },
}


def _text(value: Any) -> str:
    return str(value or "").strip()


def _parse_date(value: str | None, fallback: date) -> date:
    try:
        return date.fromisoformat((value or "")[:10])
    except ValueError:
        return fallback


def normalise_report_filters(values: Any) -> dict[str, Any]:
    today = date.today()
    report_type = _text(values.get("report_type")) or "patient_summary"
    if report_type not in REPORT_DEFINITIONS:
        report_type = "patient_summary"
    date_to = _parse_date(values.get("date_to"), today)
    date_from = _parse_date(values.get("date_from"), today - timedelta(days=30))
    if date_from > date_to:
        date_from, date_to = date_to, date_from

    def as_int(name: str) -> int | None:
        try:
            return int(values.get(name) or 0) or None
        except (TypeError, ValueError):
            return None

    risk = _text(values.get("risk")).lower() or "all"
    if risk not in {"all", "critical", "high", "high_critical", "medium", "low"}:
        risk = "all"
    return {
        "report_type": report_type,
        "date_from": date_from.isoformat(),
        "date_to": date_to.isoformat(),
        "patient_id": as_int("patient_id"),
        "risk": risk,
        "condition": _text(values.get("condition")),
        "clinician_id": as_int("clinician_id"),
        "event_type": _text(values.get("event_type")),
        "status": _text(values.get("status")),
        "source": _text(values.get("source")),
        "location": _text(values.get("location")),
    }


def report_filter_options(*, user_id: int, role: str) -> dict[str, Any]:
    ids = accessible_patient_ids(user_id=user_id, role=role, include_archived=False)
    if ids:
        marks = ",".join("?" for _ in ids)
        patients = query_db(
            f"SELECT id,external_ref,first_name,last_name FROM patients WHERE active=1 AND id IN ({marks}) ORDER BY last_name,first_name",
            ids,
        )
        event_types = query_db(
            f"""SELECT event_type value FROM care_events WHERE patient_id IN ({marks})
                UNION SELECT event_type value FROM device_events WHERE patient_id IN ({marks})
                UNION SELECT alert_type value FROM alerts WHERE patient_id IN ({marks})
                UNION SELECT scenario value FROM agentic_runs WHERE patient_id IN ({marks})
                ORDER BY value""",
            (*ids, *ids, *ids, *ids),
        )
        sources = query_db(
            f"""SELECT source value FROM vitals WHERE patient_id IN ({marks}) AND COALESCE(source,'')!=''
                UNION SELECT source value FROM care_events WHERE patient_id IN ({marks}) AND COALESCE(source,'')!=''
                UNION SELECT adapter_key value FROM device_adapter_measurements WHERE patient_id IN ({marks})
                UNION SELECT provider value FROM wearable_measurements WHERE patient_id IN ({marks})
                ORDER BY value""",
            (*ids, *ids, *ids, *ids),
        )
        statuses = query_db(
            f"""SELECT status value FROM alerts WHERE patient_id IN ({marks}) AND COALESCE(status,'')!=''
                UNION SELECT status value FROM agentic_runs WHERE patient_id IN ({marks}) AND COALESCE(status,'')!=''
                UNION SELECT status value FROM teleconsultations WHERE patient_id IN ({marks}) AND COALESCE(status,'')!=''
                UNION SELECT status value FROM notifications WHERE patient_id IN ({marks}) AND COALESCE(status,'')!=''
                UNION SELECT workflow_completion_status value FROM agentic_workflow_metrics awm
                  JOIN agentic_runs ar ON ar.id=awm.run_id WHERE ar.patient_id IN ({marks}) AND COALESCE(workflow_completion_status,'')!=''
                ORDER BY value""",
            (*ids, *ids, *ids, *ids, *ids),
        )
    else:
        patients, event_types, sources, statuses = [], [], [], []
    pop_opts = population_filter_options(user_id=user_id, role=role)
    return {
        "patients": [dict(r) for r in patients],
        "conditions": pop_opts["conditions"],
        "locations": pop_opts["locations"],
        "clinicians": pop_opts["clinicians"],
        "event_types": [_text(r["value"]) for r in event_types if _text(r["value"])],
        "sources": [_text(r["value"]) for r in sources if _text(r["value"])],
        "statuses": [_text(r["value"]) for r in statuses if _text(r["value"])],
        "excel_available": Workbook is not None,
    }


def _selected_patients(*, user_id: int, role: str, filters: dict[str, Any], force_risk: str | None = None) -> list[dict[str, Any]]:
    pop_filters = normalise_population_filters({
        "risk": force_risk or filters["risk"],
        "condition": filters["condition"],
        "location": filters["location"],
        "clinician_id": filters["clinician_id"],
        "status": "active",
        "date_from": filters["date_from"],
        "date_to": filters["date_to"],
    })
    rows = population_dashboard(user_id=user_id, role=role, filters=pop_filters)["patients"]
    if filters["patient_id"]:
        rows = [p for p in rows if int(p["id"]) == filters["patient_id"]]
    return rows


def _patient_map(rows: list[dict[str, Any]]) -> dict[int, dict[str, Any]]:
    return {int(p["id"]): p for p in rows}


def _between(column: str, filters: dict[str, Any]) -> tuple[str, tuple[str, str]]:
    return f"datetime({column}) BETWEEN datetime(?) AND datetime(?)", (f"{filters['date_from']} 00:00:00", f"{filters['date_to']} 23:59:59")


def _in_clause(ids: list[int]) -> tuple[str, tuple[int, ...]]:
    if not ids:
        return "NULL", ()
    return ",".join("?" for _ in ids), tuple(ids)


def _fmt_patient(p: dict[str, Any]) -> str:
    return f"{p.get('first_name','')} {p.get('last_name','')}".strip()


def _patient_rows(patients: list[dict[str, Any]]) -> tuple[list[dict[str, Any]], list[tuple[str, str]]]:
    rows = []
    for p in patients:
        rows.append({
            "patient_id": p.get("external_ref"),
            "name": _fmt_patient(p),
            "age": p.get("age"),
            "location": p.get("city") or "",
            "conditions": p.get("condition_summary") or "",
            "risk": p.get("priority", "low").title(),
            "risk_score": p.get("risk_score") or 0,
            "assigned_clinician": p.get("assigned_professional") or "Unassigned",
            "latest_activity": p.get("latest_activity") or "",
            "latest_vitals": p.get("latest_vitals_at") or "",
            "device_status": "Connected / data present" if p.get("wearable_connected") else "No device data",
        })
    cols = [
        ("patient_id", "Patient ID"), ("name", "Name"), ("age", "Age"), ("location", "Location"),
        ("conditions", "Conditions"), ("risk", "Risk"), ("risk_score", "Risk score"),
        ("assigned_clinician", "Assigned clinician"), ("latest_activity", "Latest alert/event"),
        ("latest_vitals", "Latest vitals"), ("device_status", "Wearable/device"),
    ]
    return rows, cols


def _query_vitals(patients: list[dict[str, Any]], filters: dict[str, Any]):
    ids = [int(p["id"]) for p in patients]
    marks, id_args = _in_clause(ids)
    if not ids:
        return [], []
    time_sql, time_args = _between("v.measured_at", filters)
    sql = f"""SELECT v.*,p.external_ref,p.first_name,p.last_name,p.city
              FROM vitals v JOIN patients p ON p.id=v.patient_id
              WHERE v.patient_id IN ({marks}) AND {time_sql}"""
    args: list[Any] = [*id_args, *time_args]
    if filters["source"]:
        sql += " AND lower(COALESCE(v.source,''))=lower(?)"
        args.append(filters["source"])
    sql += " ORDER BY v.measured_at DESC,v.id DESC"
    raw = query_db(sql, args)
    rows = [{
        "patient_id": r["external_ref"], "name": f"{r['first_name']} {r['last_name']}", "location": r["city"] or "",
        "metric": r["kind"], "value": r["value"], "unit": r["unit"] or "", "source": r["source"] or "", "measured_at": r["measured_at"],
    } for r in raw]
    cols = [("patient_id","Patient ID"),("name","Name"),("location","Location"),("metric","Metric"),("value","Value"),("unit","Unit"),("source","Source"),("measured_at","Measured at")]
    return rows, cols


def _query_agentic(patients: list[dict[str, Any]], filters: dict[str, Any]):
    ids = [int(p["id"]) for p in patients]; marks, id_args = _in_clause(ids)
    if not ids: return [], []
    time_sql, time_args = _between("ar.started_at", filters)
    sql = f"""SELECT ar.*,p.external_ref,p.first_name,p.last_name,
                     awm.response_time_seconds,awm.intervention,awm.outcome_status,awm.teleconsultation_status
              FROM agentic_runs ar JOIN patients p ON p.id=ar.patient_id
              LEFT JOIN agentic_workflow_metrics awm ON awm.run_id=ar.id
              WHERE ar.patient_id IN ({marks}) AND {time_sql}"""
    args: list[Any] = [*id_args, *time_args]
    if filters["status"]:
        sql += " AND lower(ar.status)=lower(?)"; args.append(filters["status"])
    if filters["event_type"]:
        sql += " AND lower(ar.scenario) LIKE lower(?)"; args.append(f"%{filters['event_type']}%")
    sql += " ORDER BY ar.started_at DESC"
    raw = query_db(sql, args)
    rows = [{
        "patient_id": r["external_ref"], "name": f"{r['first_name']} {r['last_name']}", "scenario": r["scenario"],
        "severity": r["severity"], "status": r["status"], "started_at": r["started_at"],
        "response_time_seconds": r["response_time_seconds"] if r["response_time_seconds"] is not None else "",
        "intervention": r["intervention"] or "", "outcome": r["outcome_status"] or "", "teleconsultation": r["teleconsultation_status"] or "",
    } for r in raw]
    cols = [("patient_id","Patient ID"),("name","Name"),("scenario","Event / scenario"),("severity","Severity"),("status","Status"),("started_at","Started"),("response_time_seconds","Response sec"),("intervention","Intervention"),("outcome","Outcome"),("teleconsultation","Teleconsultation")]
    return rows, cols


def _query_fall_deterioration(patients: list[dict[str, Any]], filters: dict[str, Any]):
    ids = [int(p["id"]) for p in patients]; marks, id_args = _in_clause(ids)
    if not ids: return [], []
    start, end = f"{filters['date_from']} 00:00:00", f"{filters['date_to']} 23:59:59"
    raw_rows: list[dict[str, Any]] = []
    ars = query_db(
        f"""SELECT ar.patient_id,p.external_ref,p.first_name,p.last_name,'Agentic Care' source_type,
                   ar.scenario event_type,ar.severity severity,ar.status status,ar.summary description,ar.started_at occurred_at
            FROM agentic_runs ar JOIN patients p ON p.id=ar.patient_id
            WHERE ar.patient_id IN ({marks}) AND ar.started_at BETWEEN ? AND ?
              AND (lower(ar.scenario) LIKE '%fall%' OR lower(ar.scenario) LIKE '%deterior%' OR lower(ar.scenario) LIKE '%declin%' OR lower(ar.scenario) LIKE '%instabil%')""",
        (*id_args, start, end),
    )
    ces = query_db(
        f"""SELECT ce.patient_id,p.external_ref,p.first_name,p.last_name,'Care event' source_type,
                   ce.event_type event_type,'' severity,'' status,ce.description description,ce.created_at occurred_at
            FROM care_events ce JOIN patients p ON p.id=ce.patient_id
            WHERE ce.patient_id IN ({marks}) AND ce.created_at BETWEEN ? AND ?
              AND (lower(ce.event_type) LIKE '%fall%' OR lower(ce.event_type) LIKE '%deterior%' OR lower(COALESCE(ce.description,'')) LIKE '%fall%' OR lower(COALESCE(ce.description,'')) LIKE '%deterior%')""",
        (*id_args, start, end),
    )
    des = query_db(
        f"""SELECT de.patient_id,p.external_ref,p.first_name,p.last_name,'Device event' source_type,
                   de.event_type event_type,de.severity severity,'' status,de.value description,de.created_at occurred_at
            FROM device_events de JOIN patients p ON p.id=de.patient_id
            WHERE de.patient_id IN ({marks}) AND de.created_at BETWEEN ? AND ?
              AND (lower(de.event_type) LIKE '%fall%' OR lower(de.event_type) LIKE '%deterior%')""",
        (*id_args, start, end),
    )
    for r in [*ars, *ces, *des]:
        item = dict(r)
        item["patient_id"] = item.pop("external_ref")
        item["name"] = f"{item.pop('first_name')} {item.pop('last_name')}"
        if filters["event_type"] and filters["event_type"].casefold() not in _text(item["event_type"]).casefold():
            continue
        if filters["status"] and _text(item["status"]).casefold() != filters["status"].casefold():
            continue
        raw_rows.append(item)
    raw_rows.sort(key=lambda r: _text(r.get("occurred_at")), reverse=True)
    cols = [("patient_id","Patient ID"),("name","Name"),("source_type","Source"),("event_type","Event"),("severity","Severity"),("status","Status"),("description","Description"),("occurred_at","Occurred at")]
    return raw_rows, cols


def _query_consultations(patients: list[dict[str, Any]], filters: dict[str, Any]):
    ids = [int(p["id"]) for p in patients]; marks, id_args = _in_clause(ids)
    if not ids: return [], []
    time_sql, time_args = _between("t.created_at", filters)
    sql = f"""SELECT t.*,p.external_ref,p.first_name,p.last_name FROM teleconsultations t
              JOIN patients p ON p.id=t.patient_id WHERE t.patient_id IN ({marks}) AND {time_sql}"""
    args: list[Any] = [*id_args, *time_args]
    if filters["status"]: sql += " AND lower(t.status)=lower(?)"; args.append(filters["status"])
    sql += " ORDER BY t.created_at DESC"
    raw = query_db(sql, args)
    rows = [{
        "patient_id": r["external_ref"], "name": f"{r['first_name']} {r['last_name']}", "clinician": r["clinician_name"] or r["clinician_role"],
        "status": r["status"], "reason": r["reason"] or "", "scheduled_at": r["scheduled_at"] or "", "started_at": r["started_at"] or "", "ended_at": r["ended_at"] or "", "created_at": r["created_at"],
    } for r in raw]
    cols = [("patient_id","Patient ID"),("name","Name"),("clinician","Clinician"),("status","Status"),("reason","Reason"),("scheduled_at","Scheduled"),("started_at","Started"),("ended_at","Ended"),("created_at","Created")]
    return rows, cols


def _query_family_notifications(patients: list[dict[str, Any]], filters: dict[str, Any]):
    ids = [int(p["id"]) for p in patients]; marks, id_args = _in_clause(ids)
    if not ids: return [], []
    time_sql, time_args = _between("n.created_at", filters)
    sql = f"""SELECT n.*,p.external_ref,p.first_name,p.last_name FROM notifications n
              JOIN patients p ON p.id=n.patient_id
              WHERE n.patient_id IN ({marks}) AND {time_sql} AND lower(COALESCE(n.audience,'')) LIKE '%family%'"""
    args: list[Any] = [*id_args, *time_args]
    if filters["status"]: sql += " AND lower(n.status)=lower(?)"; args.append(filters["status"])
    if filters["source"]: sql += " AND lower(COALESCE(n.source,''))=lower(?)"; args.append(filters["source"])
    sql += " ORDER BY n.created_at DESC"
    raw = query_db(sql, args)
    rows = [{"patient_id":r["external_ref"],"name":f"{r['first_name']} {r['last_name']}","channel":r["channel"],"status":r["status"],"source":r["source"],"message":r["message"],"created_at":r["created_at"]} for r in raw]
    cols = [("patient_id","Patient ID"),("name","Name"),("channel","Channel"),("status","Status"),("source","Source"),("message","Message"),("created_at","Created")]
    return rows, cols


def _query_device_measurements(patients: list[dict[str, Any]], filters: dict[str, Any]):
    ids = [int(p["id"]) for p in patients]; marks, id_args = _in_clause(ids)
    if not ids: return [], []
    start, end = f"{filters['date_from']} 00:00:00", f"{filters['date_to']} 23:59:59"
    rows: list[dict[str, Any]] = []
    adapter = query_db(
        f"""SELECT m.*,p.external_ref,p.first_name,p.last_name FROM device_adapter_measurements m
            JOIN patients p ON p.id=m.patient_id WHERE m.patient_id IN ({marks}) AND datetime(m.measured_at) BETWEEN datetime(?) AND datetime(?)
            ORDER BY m.measured_at DESC""", (*id_args,start,end))
    wearable = query_db(
        f"""SELECT wm.*,p.external_ref,p.first_name,p.last_name FROM wearable_measurements wm
            JOIN patients p ON p.id=wm.patient_id WHERE wm.patient_id IN ({marks}) AND datetime(wm.measured_at) BETWEEN datetime(?) AND datetime(?)
            ORDER BY wm.measured_at DESC""", (*id_args,start,end))
    for r in adapter:
        source = r["adapter_key"]
        if filters["source"] and source.casefold() != filters["source"].casefold(): continue
        value = r["value_real"] if r["value_real"] is not None else r["value_text"]
        rows.append({"patient_id":r["external_ref"],"name":f"{r['first_name']} {r['last_name']}","source":source,"metric":r["metric"],"value":value,"unit":r["unit"] or "","device":"Adapter","measured_at":r["measured_at"]})
    for r in wearable:
        source = r["provider"] or r["source"] or "wearable"
        if filters["source"] and source.casefold() != filters["source"].casefold(): continue
        rows.append({"patient_id":r["external_ref"],"name":f"{r['first_name']} {r['last_name']}","source":source,"metric":r["metric"],"value":r["value"],"unit":r["unit"] or "","device":r["device_name"] or "Wearable","measured_at":r["measured_at"]})
    rows.sort(key=lambda r: _text(r["measured_at"]), reverse=True)
    cols = [("patient_id","Patient ID"),("name","Name"),("source","Source"),("device","Device"),("metric","Metric"),("value","Value"),("unit","Unit"),("measured_at","Measured at")]
    return rows, cols


def _query_outcomes(patients: list[dict[str, Any]], filters: dict[str, Any]):
    ids = [int(p["id"]) for p in patients]; marks, id_args = _in_clause(ids)
    if not ids: return [], []
    start,end=f"{filters['date_from']} 00:00:00",f"{filters['date_to']} 23:59:59"
    raw=query_db(
        f"""SELECT ar.patient_id,p.external_ref,p.first_name,p.last_name,ar.scenario,ar.status run_status,
                   awm.response_time_seconds,awm.intervention,awm.outcome_status,awm.workflow_completion_status,
                   awm.teleconsultation_status,awm.operational_impact,awm.updated_at
            FROM agentic_workflow_metrics awm JOIN agentic_runs ar ON ar.id=awm.run_id JOIN patients p ON p.id=ar.patient_id
            WHERE ar.patient_id IN ({marks}) AND COALESCE(awm.event_detected_at,ar.started_at) BETWEEN ? AND ?
            ORDER BY COALESCE(awm.event_detected_at,ar.started_at) DESC""",(*id_args,start,end))
    rows=[]
    for r in raw:
        if filters["status"] and _text(r["workflow_completion_status"]).casefold()!=filters["status"].casefold(): continue
        rows.append({"patient_id":r["external_ref"],"name":f"{r['first_name']} {r['last_name']}","scenario":r["scenario"],"response_time_seconds":r["response_time_seconds"] if r["response_time_seconds"] is not None else "","intervention":r["intervention"] or "","outcome":r["outcome_status"] or "","workflow_status":r["workflow_completion_status"] or r["run_status"],"teleconsultation":r["teleconsultation_status"] or "","operational_impact":r["operational_impact"] or "","updated_at":r["updated_at"]})
    cols=[("patient_id","Patient ID"),("name","Name"),("scenario","Scenario"),("response_time_seconds","Response sec"),("intervention","Intervention"),("outcome","Outcome"),("workflow_status","Workflow status"),("teleconsultation","Teleconsultation"),("operational_impact","Operational impact"),("updated_at","Updated")]
    return rows,cols


def _query_alerts(patients: list[dict[str, Any]], filters: dict[str, Any]):
    ids=[int(p["id"]) for p in patients]; marks,id_args=_in_clause(ids)
    if not ids:return [],[]
    time_sql,time_args=_between("a.created_at",filters)
    sql=f"""SELECT a.*,p.external_ref,p.first_name,p.last_name FROM alerts a JOIN patients p ON p.id=a.patient_id
             WHERE a.patient_id IN ({marks}) AND {time_sql}"""
    args:list[Any]=[*id_args,*time_args]
    if filters["status"]:sql+=" AND lower(a.status)=lower(?)";args.append(filters["status"])
    if filters["event_type"]:sql+=" AND lower(a.alert_type)=lower(?)";args.append(filters["event_type"])
    sql+=" ORDER BY a.created_at DESC"
    raw=query_db(sql,args)
    rows=[{"patient_id":r["external_ref"],"name":f"{r['first_name']} {r['last_name']}","alert_type":r["alert_type"],"severity":r["severity"],"status":r["status"],"message":r["message"],"created_at":r["created_at"]} for r in raw]
    cols=[("patient_id","Patient ID"),("name","Name"),("alert_type","Alert type"),("severity","Severity"),("status","Status"),("message","Message"),("created_at","Created")]
    return rows,cols


def _population_summary(*, user_id:int,role:str,filters:dict[str,Any]):
    pop_filters=normalise_population_filters({"risk":filters["risk"],"condition":filters["condition"],"location":filters["location"],"clinician_id":filters["clinician_id"],"date_from":filters["date_from"],"date_to":filters["date_to"],"status":"active"})
    data=population_dashboard(user_id=user_id,role=role,filters=pop_filters)
    rows=[{"metric":"Selected population","value":data["metrics"]["total"]}]
    labels={"high_critical":"High / Critical risk","medium":"Medium risk","low":"Low risk","fall_risk":"Fall risk","copd":"COPD","hypertension":"Hypertension","diabetes":"Diabetes","wearable_recent":"Recent wearable/device data","recent_alerts":"Patients with recent alerts","recent_consultations":"Patients with recent consultations","open_agentic":"Open Agentic Care workflows","open_alerts":"Open alerts"}
    for key,label in labels.items(): rows.append({"metric":label,"value":data["metrics"].get(key,0)})
    return rows,[("metric","Metric"),("value","Value")]


def generate_report(*,user_id:int,role:str,generated_by:str,filters:dict[str,Any])->dict[str,Any]:
    report_type=filters["report_type"]
    force_risk="high_critical" if report_type=="high_critical_risk" else None
    patients=_selected_patients(user_id=user_id,role=role,filters=filters,force_risk=force_risk)
    builders:dict[str,Callable]= {
        "patient_summary":lambda:_patient_rows(patients),
        "high_critical_risk":lambda:_patient_rows(patients),
        "vitals_summary":lambda:_query_vitals(patients,filters),
        "agentic_events":lambda:_query_agentic(patients,filters),
        "fall_deterioration":lambda:_query_fall_deterioration(patients,filters),
        "consultations":lambda:_query_consultations(patients,filters),
        "family_notifications":lambda:_query_family_notifications(patients,filters),
        "device_measurements":lambda:_query_device_measurements(patients,filters),
        "outcome_metrics":lambda:_query_outcomes(patients,filters),
        "clinical_alerts":lambda:_query_alerts(patients,filters),
        "population_summary":lambda:_population_summary(user_id=user_id,role=role,filters=filters),
    }
    rows,columns=builders[report_type]()
    definition=REPORT_DEFINITIONS[report_type]
    patient_count = len({r.get("patient_id") for r in rows if r.get("patient_id")})
    summary=[{"label":"Rows","value":len(rows)},{"label":"Patients","value":patient_count or len(patients)},{"label":"Date from","value":filters["date_from"]},{"label":"Date to","value":filters["date_to"]}]
    if report_type in {"patient_summary","high_critical_risk"}:
        summary=[{"label":"Patients","value":len(rows)},{"label":"High / Critical","value":sum(1 for r in rows if r.get("risk") in {"High","Critical"})}]
    elif report_type == "vitals_summary":
        summary=[{"label":"Measurements","value":len(rows)},{"label":"Patients","value":patient_count},{"label":"Metric types","value":len({r.get('metric') for r in rows})}]
    elif report_type == "agentic_events":
        summary=[{"label":"Agentic events","value":len(rows)},{"label":"Patients","value":patient_count},{"label":"High / Critical","value":sum(1 for r in rows if _text(r.get('severity')).lower() in {'high','critical'})}]
    elif report_type == "fall_deterioration":
        summary=[{"label":"Events","value":len(rows)},{"label":"Patients","value":patient_count},{"label":"Fall events","value":sum(1 for r in rows if 'fall' in _text(r.get('event_type')).casefold())}]
    elif report_type == "consultations":
        summary=[{"label":"Consultations","value":len(rows)},{"label":"Patients","value":patient_count},{"label":"Open / requested","value":sum(1 for r in rows if _text(r.get('status')).lower() in {'open','requested','scheduled'})}]
    elif report_type == "family_notifications":
        summary=[{"label":"Family notifications","value":len(rows)},{"label":"Patients","value":patient_count}]
    elif report_type == "device_measurements":
        summary=[{"label":"Measurements","value":len(rows)},{"label":"Patients","value":patient_count},{"label":"Sources","value":len({r.get('source') for r in rows})}]
    elif report_type == "outcome_metrics":
        response_values=[float(r['response_time_seconds']) for r in rows if str(r.get('response_time_seconds') or '').replace('.','',1).isdigit()]
        summary=[{"label":"Workflows","value":len(rows)},{"label":"Patients","value":patient_count},{"label":"Avg response sec","value":round(sum(response_values)/len(response_values),1) if response_values else '—'}]
    elif report_type == "clinical_alerts":
        summary=[{"label":"Alerts","value":len(rows)},{"label":"Open","value":sum(1 for r in rows if _text(r.get('status')).lower()=='open')},{"label":"High / Critical","value":sum(1 for r in rows if _text(r.get('severity')).lower() in {'high','critical'})}]
    elif report_type == "population_summary":
        summary=[{"label":"Population metrics","value":len(rows)},{"label":"Patients in scope","value":len(patients)}]
    return {
        "report_type":report_type,"title":definition["label"],"description":definition["description"],
        "columns":[{"key":k,"label":label} for k,label in columns],"rows":rows,"summary":summary,
        "filters":filters.copy(),"generated_at":datetime.now().astimezone().strftime("%Y-%m-%d %H:%M:%S %Z"),"generated_by":generated_by,
        "empty_message":"No records match the selected filters. CareAI has not generated placeholder data.",
    }


def report_filename(result:dict[str,Any],extension:str)->str:
    stamp=datetime.now().strftime("%Y%m%d_%H%M%S")
    return f"careai_{result['report_type']}_{stamp}.{extension}"


def _safe_export_value(value: Any) -> Any:
    """Prevent spreadsheet formula execution from text fields."""
    if isinstance(value, str) and value[:1] in {"=", "+", "-", "@"}:
        return "'" + value
    return value


def build_csv(result:dict[str,Any])->bytes:
    buf=io.StringIO(newline="")
    writer=csv.writer(buf)
    writer.writerow([result["title"]]);writer.writerow(["Generated",result["generated_at"]]);writer.writerow(["Generated by",result["generated_by"]])
    writer.writerow([]); writer.writerow(["Selected filters"])
    for key,value in result["filters"].items():
        if key != "report_type" and value not in (None,"","all"):
            writer.writerow([key.replace("_"," ").title(),value])
    writer.writerow([]); writer.writerow(["Summary"])
    for item in result["summary"]: writer.writerow([item["label"],item["value"]])
    writer.writerow([]); writer.writerow([c["label"] for c in result["columns"]])
    for row in result["rows"]:writer.writerow([_safe_export_value(row.get(c["key"],"")) for c in result["columns"]])
    return buf.getvalue().encode("utf-8-sig")


def build_excel(result:dict[str,Any])->bytes:
    if Workbook is None: raise RuntimeError("Excel export is not available because openpyxl is not installed.")
    wb=Workbook();ws=wb.active;ws.title="Report"
    ws.append([result["title"]]);ws.append(["Generated",result["generated_at"]]);ws.append(["Generated by",result["generated_by"]]);ws.append([])
    ws.append(["Selected filters"])
    for key,value in result["filters"].items():
        if key != "report_type" and value not in (None,"","all"):
            ws.append([key.replace("_"," ").title(),value])
    ws.append([]); ws.append(["Summary"])
    for item in result["summary"]: ws.append([item["label"],item["value"]])
    ws.append([])
    header_row=ws.max_row+1
    ws.append([c["label"] for c in result["columns"]])
    for row in result["rows"]:ws.append([_safe_export_value(row.get(c["key"],"")) for c in result["columns"]])
    ws.freeze_panes=f"A{header_row+1}"
    if result["rows"]:
        ws.auto_filter.ref=f"A{header_row}:{ws.cell(ws.max_row, len(result['columns'])).coordinate}"
    for column_cells in ws.columns:
        max_len=max((len(str(cell.value or "")) for cell in column_cells),default=10)
        ws.column_dimensions[column_cells[0].column_letter].width=min(max(max_len+2,10),45)
    out=io.BytesIO();wb.save(out);return out.getvalue()


def build_pdf(result:dict[str,Any])->bytes:
    out=io.BytesIO();doc=SimpleDocTemplate(out,pagesize=landscape(A4),rightMargin=10*mm,leftMargin=10*mm,topMargin=10*mm,bottomMargin=10*mm)
    styles=getSampleStyleSheet();small=styles["BodyText"].clone("CareAISmall");small.fontSize=6;small.leading=7
    story=[Paragraph(escape(_text(result["title"])),styles["Title"]),Paragraph(escape(f"Generated {result['generated_at']} by {result['generated_by']}"),styles["Normal"]),Spacer(1,4*mm)]
    active_filters=[f"{k.replace('_',' ').title()}: {v}" for k,v in result["filters"].items() if v not in (None,"","all") and k!="report_type"]
    if active_filters:story.extend([Paragraph(escape("Filters: "+" | ".join(active_filters)),small),Spacer(1,3*mm)])
    summary_text=" | ".join(f"{item['label']}: {item['value']}" for item in result["summary"])
    if summary_text: story.extend([Paragraph(escape("Summary: "+summary_text),small),Spacer(1,3*mm)])
    if not result["rows"]:
        story.append(Paragraph(escape(_text(result["empty_message"])),styles["Normal"]));doc.build(story);return out.getvalue()
    cols=result["columns"]
    max_pdf_rows=500
    data=[[Paragraph(escape(str(c["label"])),small) for c in cols]]
    for row in result["rows"][:max_pdf_rows]:
        data.append([Paragraph(escape(_text(row.get(c["key"],""))[:500]),small) for c in cols])
    table=Table(data,repeatRows=1)
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#e9f1ed")),("TEXTCOLOR",(0,0),(-1,0),colors.HexColor("#164c3a")),("GRID",(0,0),(-1,-1),0.25,colors.HexColor("#c7d2cc")),("VALIGN",(0,0),(-1,-1),"TOP"),("FONTSIZE",(0,0),(-1,-1),6),("LEFTPADDING",(0,0),(-1,-1),3),("RIGHTPADDING",(0,0),(-1,-1),3)]))
    story.append(table)
    if len(result["rows"])>max_pdf_rows:story.append(Paragraph(f"PDF preview limited to first {max_pdf_rows} rows. CSV/Excel contains all {len(result['rows'])} rows.",small))
    doc.build(story);return out.getvalue()


def recent_report_activity(*,user_id:int,limit:int=8)->list[dict[str,Any]]:
    rows=query_db("""SELECT action,details,created_at FROM audit_logs WHERE user_id=? AND action IN ('GENERATE_REPORT','DOWNLOAD_REPORT') ORDER BY id DESC LIMIT ?""",(user_id,limit))
    result=[]
    for row in rows:
        item=dict(row)
        try:
            detail=json.loads(item.get("details") or "{}")
            report_type=detail.get("report_type") or "report"
            label=REPORT_DEFINITIONS.get(report_type,{}).get("label",report_type.replace("_"," ").title())
            suffix=f" · {detail.get('format','').upper()}" if detail.get("format") else ""
            item["details"]=f"{label}{suffix} · {int(detail.get('row_count') or 0)} row(s)"
        except (ValueError,TypeError,AttributeError):
            pass
        result.append(item)
    return result
